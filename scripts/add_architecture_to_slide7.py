from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "dist" / "go-ing-fast-with-mcp.pptx"


def set_text(shape, text, size=14, bold=False, color="191919"):
    shape.text_frame.clear()
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = shape.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_box(slide, left, top, width, height, text, fill, line="4472C4"):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    shape.line.color.rgb = RGBColor.from_string(line)
    shape.line.width = Pt(1.25)
    set_text(shape, text)
    return shape


def add_arrow(slide, x, y1, y2):
    arrow = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x), Inches(y1), Inches(x), Inches(y2),
    )
    arrow.line.color.rgb = RGBColor.from_string("5F5F5F")
    arrow.line.width = Pt(1.25)
    arrow.line.end_arrowhead = True
    return arrow


def main():
    prs = Presentation(DECK)
    slide = prs.slides[6]  # Slide 7: A Go template

    # The existing slide has open space beneath the Code Structure label.
    left, width = 1.35, 7.55
    box_left, box_width, box_height = 1.65, 6.95, 0.48
    ys = [3.02, 3.72, 4.42, 5.12]
    labels = [
        "MCP client / agent",
        "MCP adapter  ·  schemas, tools, protocol errors",
        "Shared core  ·  auth, client, validation, limits",
        "Vendor API / SDK / gRPC  or  local CLI",
    ]
    fills = ["EAF2F8", "D9EAF7", "E2F0D9", "F2F2F2"]

    heading = slide.shapes.add_textbox(
        Inches(left), Inches(2.52), Inches(width), Inches(0.28)
    )
    set_text(heading, "Architecture pattern: keep MCP thin", size=13, bold=True)
    heading.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT

    for y, label, fill in zip(ys, labels, fills):
        add_box(slide, box_left, y, box_width, box_height, label, fill)

    for y1, y2 in zip(ys[:-1], ys[1:]):
        add_arrow(slide, box_left + box_width / 2, y1 + box_height, y2)

    note = slide.shapes.add_textbox(
        Inches(8.85), Inches(3.10), Inches(1.25), Inches(1.65)
    )
    note.text_frame.clear()
    note.text_frame.word_wrap = True
    p = note.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = "Reuse the core\nin the CLI, tests,\nand MCP server."
    run.font.name = "Aptos"
    run.font.size = Pt(11)
    run.font.italic = True
    run.font.color.rgb = RGBColor.from_string("5F5F5F")

    prs.save(DECK)


if __name__ == "__main__":
    main()
