"""Generate a deliberately plain PowerPoint deck from slide-content.md notes."""

from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


OUT = Path("dist/go-ing-fast-with-mcp.pptx")
W, H = Inches(13.333), Inches(7.5)
BLACK = RGBColor(25, 25, 25)
GRAY = RGBColor(95, 95, 95)
LIGHT = RGBColor(225, 225, 225)


def add_text(slide, text, x, y, w, h, size=24, bold=False, color=BLACK,
             font="Aptos", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def base_slide(prs, title, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(slide, title, Inches(0.75), Inches(0.45), Inches(11.8), Inches(0.55), 28, True)
    line = slide.shapes.add_shape(1, Inches(0.75), Inches(1.12), Inches(11.85), Inches(0.012))
    line.fill.solid()
    line.fill.fore_color.rgb = LIGHT
    line.line.fill.background()
    add_text(slide, str(number), Inches(12.35), Inches(7.02), Inches(0.35), Inches(0.2), 9, color=GRAY,
             align=PP_ALIGN.RIGHT)
    return slide


def bullets(slide, items, x=0.95, y=1.5, w=11.3, h=5.1, size=22):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = BLACK
        p.space_after = Pt(15)
        p.bullet = True
    return box


def quote(slide, text, y=6.18):
    add_text(slide, text, Inches(0.95), Inches(y), Inches(11.35), Inches(0.5), 15, color=GRAY)


def code(slide, text, y=1.55, h=4.75):
    box = slide.shapes.add_textbox(Inches(1.1), Inches(y), Inches(11.0), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(247, 247, 247)
    box.line.color.rgb = LIGHT
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Aptos Mono"
    p.font.size = Pt(18)
    p.font.color.rgb = BLACK
    return box


def main():
    OUT.parent.mkdir(exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    prs.core_properties.title = "Go-ing Fast with MCP"
    prs.core_properties.subject = "DC State of the Stack 2026"
    prs.core_properties.author = "Matthew Franz"

    # Cover
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(s, "Go-ing Fast with MCP", Inches(0.9), Inches(1.75), Inches(11.5), Inches(0.8), 40, True)
    add_text(s, "Rapid Development of MCP Servers for Security Agents", Inches(0.92), Inches(2.7), Inches(10.8), Inches(0.5), 24, color=GRAY)
    add_text(s, "Matthew Franz  |  DC State of the Stack 2026", Inches(0.92), Inches(5.85), Inches(10.8), Inches(0.35), 17, color=GRAY)

    slides = [
        ("During the next five minutes…", [
            "Why you should consider writing a custom MCP server for data-intensive platforms (SIEM, observability)",
            "How to iteratively build, test, and run STDIO MCP servers — for Python security agents or Claude, Codex, OpenCode",
        ], None),
        ("Who am I?", [
            "Security & Compliance Lead, Pydantic — Logfire, plus security agents with Pydantic AI",
            "25 years — Cisco, Tenable, Mandiant, Ping Identity, T. Rowe Price",
            "Last 2.5 years — building blue-team security agents",
            "This talk — lessons from CrowdStrike, Chronicle, Coralogix, and more",
        ], "Built a dedicated dozen MCP servers in Go, Rust, or Python in 2026."),
        ("Why MCP?", [
            "Token burn — large JSON payloads overwhelming agent context",
            "Agent thrashing — hallucinated queries and API calls leading to agent delays",
            "Security — inherent vulnerabilities, RBAC challenges, and more",
        ], "“USB for agents,” with well-understood challenges. Not all-or-nothing: MCP works alongside CLI tools and CodeMode."),
        ("Why custom MCP?", [
            "No logging or caching — black boxes",
            "Generic tool calls — tuned for everyone, not your telemetry",
            "Sketchy MCP marketplaces — do you really trust them with your API keys?",
        ], "You have to learn the API and the data anyway — package that knowledge in the server."),
        ("Why Go?", [
            "Sure, Python/JS libraries are great — but their dependencies, supply-chain risk, and startup time",
            "Single executable — CLI/TUI and MCP server",
            "Easier lift than Rust — for Python developers, easier/faster agentic coding",
        ], "Use other languages for time-expediency, or when a trusted SDK is available in Python."),
        ("Lessons — start with the CLI", [
            "Build the CLI first — reuse a client library, or generate from Swagger/OpenAPI (worst case: HTML/PDF)",
            "Explore the data — list / get / describe",
            "Start with 3–5 tools — enough to learn the shape and the use cases",
        ], None),
        ("Lessons — instrument and iterate", [
            "Instrument from day one — logging, observability, test coverage",
            "Add an LLM harness — Pydantic AI against common queries",
            "Explore, then specialize — let telemetry tell you the next tool to build",
        ], None),
        ("A Go template", ["Official MCP Go SDK", "Cobra + Viper — spf13"], "See GO-MCP-GUIDE.md"),
        ("Testing layers", [
            "Unit tests — CLI and MCP server",
            "JSON-RPC smoke test — no LLM required",
            "Pydantic AI harness — real tools, real model",
            "Interactive — explore with Claude, Codex, OpenCode, etc.",
        ], None),
        ("Start with 3 primitives", [
            "list_tables — what data exists?",
            "describe_table — which columns can I query?",
            "run_query — execute, return bounded JSON",
        ], "Simple example with OSquery. Proves protocol init, execution, and discovery — and lets you compare models."),
    ]
    for idx, (title, items, note) in enumerate(slides, start=2):
        s = base_slide(prs, title, idx)
        bullets(s, items)
        if note:
            quote(s, note)

    # Tool ladder table
    s = base_slide(prs, "The tool ladder", 12)
    rows = [
        ("See the surface", "list_tables", "Cheapest inventory"),
        ("Find relevant data", "search_tables", "Names + columns"),
        ("Inspect one schema", "describe_table", "No row payloads"),
        ("Learn shape + values", "preview_table", "Schema + sample"),
        ("Query one table", "query_table", "Validated, bounded"),
        ("Join or advanced SQL", "run_query", "Flexible, riskiest"),
    ]
    table = s.shapes.add_table(7, 3, Inches(0.95), Inches(1.55), Inches(11.35), Inches(4.45)).table
    for col, width in zip(table.columns, [3.3, 3.0, 5.05]): col.width = Inches(width)
    for c, text in enumerate(["Intent", "Tool", "Why"]):
        cell = table.cell(0, c); cell.text = text; cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(240, 240, 240)
    for r, row in enumerate(rows, start=1):
        for c, text in enumerate(row): table.cell(r, c).text = text
    for row in table.rows:
        for cell in row.cells:
            cell.margin_left = Inches(0.12)
            for p in cell.text_frame.paragraphs:
                p.font.name = "Aptos"; p.font.size = Pt(15); p.font.color.rgb = BLACK
    quote(s, "Prefer the narrowest tool that answers the question.")

    # Code slides
    s = base_slide(prs, "Tool typing — args", 13)
    code(s, '''type SearchProcessesArgs struct {
  Executable  string `jsonschema:"Exact path match. No wildcards."`
  CommandLine string `jsonschema:"Token search (OR), not substring."`
  HashSHA256  string `jsonschema:"SHA256 for malware detection."`
  // + ProcessName, Host, Start/End
}''')
    s = base_slide(prs, "Tool typing — registration", 14)
    code(s, '''mcp.AddTool(server, &mcp.Tool{
  Name: "search_processes",
  Description: "Search process events by executable,\n    command line, hash, host, time range.",
})''')
    quote(s, "The tags are the docs the model reads — no separate prompt to keep in sync.")

    tail = [
        ("Guardrails — bound the query", ["Require a filter — and always bound the time range", "Read-only identity — restrict index and scope", "Per-tool timeouts — surface partial results"], None),
        ("Guardrails — bound the response", ["Cap characters — return truncation metadata", "Cache by freshness — different TTLs per tool", "Export separately — paginate to files, never bulk JSON into context"], "Search is for reasoning. Export is for volume."),
        ("Pydantic E2E testing", ["One shared client engine — CLI, TUI, and MCP; test deterministically first", "Normalize model queries — catch hallucinated syntax, return corrective hints", "Query APIs aren't read-only — confirmed the hard way with DuckDB"], None),
        ("Implement observability early", ["Logs — one agent builds, another observes and gives feedback", "Traces — exceptions and spans", "Metrics — query latency"], "Don't roll your own: Logfire (SaaS) or Agentsview (local)."),
        ("Agentsview", ["Local agentic analytics for Claude, Codex, and OpenCode", "Tool sequence, arguments, and outcomes per task"], "https://github.com/kenn-io/agentsview"),
        ("Logfire", ["OpenTelemetry traces across model calls, tool calls, and spans", "Queryable after the fact, with a built-in MCP server"], "https://pydantic.dev/logfire"),
        ("For more info", ["osqueryi-mcp — starter server, smoke harnesses, framework examples", "elastic-security-mcp — structured tools, raw DSL fallback, caching, export", "mcp-threat-hunt-aug-2026 — the long-form companion talk"], "github.com/mdfranz — for another SIEM, bring its OpenAPI spec and rebuild the same ladder."),
    ]
    for idx, (title, items, note) in enumerate(tail, start=15):
        s = base_slide(prs, title, idx)
        bullets(s, items, size=21)
        if note:
            quote(s, note)

    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
